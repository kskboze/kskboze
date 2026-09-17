"""
CSV / Excel ファイルを読み込んでデータベースに入れる部分です。

【設計のねらい】
予約システムから出力されるCSVは、ゴルフ場や業者によって列名がバラバラです。
そこで「日本語の列名 → システム内部の項目名」の対応表を持たせ、
多少表記が違っても自動で読み取れるようにしています。
対応表に無い列名が来た場合は、取り込み結果の画面に「読めなかった列」として表示します。
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from app.db import session, upsert

# ---------------------------------------------------------------
# 列名の対応表（左：内部の項目名 / 右：CSVに出てきそうな見出しの候補）
# ---------------------------------------------------------------
COLUMN_ALIASES: dict[str, dict[str, list[str]]] = {
    "courses": {
        "course_id": ["コースID", "ゴルフ場ID", "コースコード", "施設コード", "course_id", "id"],
        "name": ["コース名", "ゴルフ場名", "施設名", "name"],
        "short_name": ["略称", "短縮名", "short_name"],
        "area": ["エリア", "地区", "都道府県", "area"],
        "capacity_weekday": ["平日定員", "平日キャパ", "平日販売枠", "capacity_weekday"],
        "capacity_weekend": ["土日定員", "休日定員", "土日キャパ", "capacity_weekend"],
        "note": ["備考", "メモ", "note"],
    },
    "customers": {
        "customer_id": ["顧客ID", "会員番号", "顧客番号", "customer_id", "id"],
        "name": ["氏名", "名前", "顧客名", "会員名", "name"],
        "kana": ["カナ", "フリガナ", "ふりがな", "kana"],
        "email": ["メールアドレス", "メール", "Eメール", "email", "mail"],
        "phone": ["電話番号", "TEL", "電話", "phone", "tel"],
        "line_id": ["LINEID", "LINE_ID", "LINE", "line_id"],
        "postal_code": ["郵便番号", "〒", "postal_code", "zip"],
        "prefecture": ["都道府県", "県", "prefecture", "pref"],
        "city": ["市区町村", "住所", "city"],
        "birth_year": ["生年", "生まれ年", "birth_year"],
        "gender": ["性別", "gender", "sex"],
        "member_type": ["会員区分", "顧客区分", "区分", "member_type"],
        "home_course_id": ["ホームコース", "主利用コース", "home_course_id"],
        "allow_mail": ["メール可否", "メール許諾", "メール配信可", "allow_mail"],
        "allow_dm": ["DM可否", "DM許諾", "DM送付可", "allow_dm"],
        "allow_line": ["LINE可否", "LINE許諾", "allow_line"],
        "allow_sms": ["SMS可否", "SMS許諾", "allow_sms"],
        "registered_at": ["登録日", "入会日", "registered_at"],
    },
    "visits": {
        "visit_id": ["来場ID", "実績ID", "伝票番号", "visit_id"],
        "customer_id": ["顧客ID", "会員番号", "顧客番号", "customer_id"],
        "course_id": ["コースID", "ゴルフ場ID", "コースコード", "course_id"],
        "play_date": ["来場日", "プレー日", "利用日", "play_date", "date"],
        "booked_at": ["予約日", "受付日", "申込日", "booked_at"],
        "players": ["人数", "来場人数", "プレー人数", "players"],
        "revenue": ["売上", "売上金額", "利用金額", "金額", "revenue"],
        "plan_name": ["プラン名", "プラン", "コース名称", "plan_name"],
        "channel": ["予約経路", "販売経路", "チャネル", "媒体", "channel"],
    },
    "reservations": {
        "reservation_id": ["予約ID", "予約番号", "reservation_id"],
        "customer_id": ["顧客ID", "会員番号", "顧客番号", "customer_id"],
        "course_id": ["コースID", "ゴルフ場ID", "コースコード", "course_id"],
        "play_date": ["来場日", "プレー日", "利用日", "play_date", "date"],
        "booked_at": ["予約日", "受付日", "申込日", "booked_at"],
        "players": ["人数", "予約人数", "players"],
        "amount": ["金額", "予約金額", "売上見込", "amount"],
        "plan_name": ["プラン名", "プラン", "plan_name"],
        "channel": ["予約経路", "販売経路", "チャネル", "媒体", "channel"],
        "status": ["状態", "ステータス", "status"],
    },
    "budgets": {
        "course_id": ["コースID", "ゴルフ場ID", "コースコード", "course_id"],
        "year_month": ["年月", "対象月", "予算月", "year_month", "month"],
        "target_revenue": ["予算", "売上予算", "目標売上", "target_revenue"],
        "target_players": ["目標人数", "人数予算", "来場目標", "target_players"],
    },
}

PRIMARY_KEYS = {
    "courses": ["course_id"],
    "customers": ["customer_id"],
    "visits": ["visit_id"],
    "reservations": ["reservation_id"],
    "budgets": ["course_id", "year_month"],
}

REQUIRED = {
    "courses": ["course_id", "name"],
    "customers": ["customer_id"],
    "visits": ["customer_id", "course_id", "play_date"],
    "reservations": ["course_id", "play_date"],
    "budgets": ["course_id", "year_month", "target_revenue"],
}

INT_FIELDS = {
    "capacity_weekday", "capacity_weekend", "birth_year", "players", "revenue",
    "amount", "target_revenue", "target_players",
    "allow_mail", "allow_dm", "allow_line", "allow_sms",
}
DATE_FIELDS = {"play_date", "booked_at", "registered_at"}


@dataclass
class ImportResult:
    """取り込みの結果をまとめたもの。画面にそのまま表示します。"""
    kind: str
    filename: str
    rows_ok: int = 0
    rows_ng: int = 0
    errors: list[str] = field(default_factory=list)
    unmapped_columns: list[str] = field(default_factory=list)
    mapped_columns: dict[str, str] = field(default_factory=dict)

    @property
    def ok(self) -> bool:
        return self.rows_ok > 0

    def summary(self) -> str:
        s = f"{self.rows_ok}件を取り込みました。"
        if self.rows_ng:
            s += f" {self.rows_ng}件は読み取れませんでした。"
        return s


# ---------------------------------------------------------------
# 値をきれいにする小さな道具たち
# ---------------------------------------------------------------
def normalize_header(text: str) -> str:
    """全角・半角や記号の違いを吸収して、見出しを比較しやすい形にします。"""
    if text is None:
        return ""
    t = unicodedata.normalize("NFKC", str(text)).strip()
    t = re.sub(r"[\s_\-・／/()（）\[\]【】\.]", "", t)
    return t.lower()


def parse_date(value: Any) -> str | None:
    """いろいろな日付表記を YYYY-MM-DD に揃えます。"""
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")
    s = unicodedata.normalize("NFKC", str(value)).strip()
    if not s:
        return None
    s = re.sub(r"[年月]", "/", s).replace("日", "")
    s = s.split()[0]
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d", "%y/%m/%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def parse_year_month(value: Any) -> str | None:
    """2026-10 / 2026年10月 / 202610 などを YYYY-MM に揃えます。"""
    if value in (None, ""):
        return None
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m")
    s = unicodedata.normalize("NFKC", str(value)).strip()
    s = re.sub(r"[年月]", "-", s).rstrip("-")
    m = re.match(r"^(\d{4})[-/\.]?(\d{1,2})", s)
    if m:
        return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}"
    return None


def parse_int(value: Any) -> int | None:
    """「1,200円」「3名」のような表記から数字だけを取り出します。"""
    if value in (None, ""):
        return None
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        return int(value)
    s = unicodedata.normalize("NFKC", str(value)).strip()
    if s in ("○", "◯", "可", "OK", "ok", "TRUE", "True", "true", "有", "はい"):
        return 1
    if s in ("×", "✕", "不可", "NG", "ng", "FALSE", "False", "false", "無", "いいえ"):
        return 0
    s = re.sub(r"[^\d\-]", "", s)
    if s in ("", "-"):
        return None
    try:
        return int(s)
    except ValueError:
        return None


# ---------------------------------------------------------------
# ファイルの読み込み
# ---------------------------------------------------------------
def read_table(content: bytes, filename: str) -> list[dict[str, Any]]:
    """CSV でも Excel でも、同じ「辞書の一覧」の形にして返します。"""
    suffix = Path(filename).suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return _read_excel(content)
    return _read_csv(content)


def _read_csv(content: bytes) -> list[dict[str, Any]]:
    # 日本の業務CSVは Shift_JIS(cp932) のことが多いので、順番に試します。
    text = None
    for encoding in ("utf-8-sig", "cp932", "utf-8", "euc_jp"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = content.decode("utf-8", errors="replace")
    sample = text[:4096]
    delimiter = "\t" if sample.count("\t") > sample.count(",") else ","
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    return [dict(row) for row in reader]


def _read_excel(content: bytes) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError(
            "Excelを読むには openpyxl が必要です。`pip install openpyxl` を実行してください。"
        ) from exc
    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    try:
        header = [str(h).strip() if h is not None else "" for h in next(rows)]
    except StopIteration:
        return []
    out = []
    for row in rows:
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        out.append({header[i]: row[i] for i in range(min(len(header), len(row)))})
    return out


def build_mapping(headers: list[str], kind: str) -> tuple[dict[str, str], list[str]]:
    """CSVの見出しを、システムの項目名に対応づけます。"""
    aliases = COLUMN_ALIASES[kind]
    lookup: dict[str, str] = {}
    for field_name, candidates in aliases.items():
        for cand in [field_name] + candidates:
            lookup[normalize_header(cand)] = field_name

    mapping: dict[str, str] = {}
    unmapped: list[str] = []
    for h in headers:
        key = normalize_header(h)
        if key in lookup and lookup[key] not in mapping.values():
            mapping[h] = lookup[key]
        elif h:
            unmapped.append(h)
    return mapping, unmapped


def clean_value(field_name: str, value: Any) -> Any:
    if field_name in DATE_FIELDS:
        return parse_date(value)
    if field_name == "year_month":
        return parse_year_month(value)
    if field_name in INT_FIELDS:
        return parse_int(value)
    if value is None:
        return None
    s = str(value).strip()
    return s or None


def import_rows(kind: str, rows: list[dict[str, Any]], filename: str = "") -> ImportResult:
    """読み込んだ行をデータベースに保存します。"""
    result = ImportResult(kind=kind, filename=filename)
    if not rows:
        result.errors.append("データが1行もありませんでした。ファイルの中身をご確認ください。")
        return result

    headers = list(rows[0].keys())
    mapping, unmapped = build_mapping(headers, kind)
    result.mapped_columns = mapping
    result.unmapped_columns = unmapped

    missing = [f for f in REQUIRED[kind] if f not in mapping.values()]
    if missing:
        labels = {f: COLUMN_ALIASES[kind][f][0] for f in missing}
        result.errors.append(
            "必須の列が見つかりません: " + "、".join(labels.values())
            + "。見出しの行が1行目にあるかご確認ください。"
        )
        return result

    cleaned: list[dict[str, Any]] = []
    for index, raw in enumerate(rows, start=2):  # 2行目からがデータ
        record: dict[str, Any] = {}
        for header, field_name in mapping.items():
            record[field_name] = clean_value(field_name, raw.get(header))

        bad = [f for f in REQUIRED[kind] if record.get(f) in (None, "")]
        if bad:
            result.rows_ng += 1
            if len(result.errors) < 10:
                result.errors.append(f"{index}行目: {', '.join(bad)} が空のため取り込めませんでした。")
            continue

        # ID列が無いCSVでも扱えるよう、自動で採番します。
        if kind == "visits" and not record.get("visit_id"):
            record["visit_id"] = f"{record['customer_id']}_{record['course_id']}_{record['play_date']}"
        if kind == "reservations" and not record.get("reservation_id"):
            record["reservation_id"] = (
                f"{record['course_id']}_{record['play_date']}_"
                f"{record.get('customer_id') or 'guest'}_{index}"
            )
        cleaned.append(record)

    if cleaned:
        with session() as conn:
            _ensure_referenced_rows(conn, kind, cleaned)
            result.rows_ok = upsert(conn, kind, cleaned, PRIMARY_KEYS[kind])
            conn.execute(
                "INSERT INTO import_log(imported_at, kind, filename, rows_ok, rows_ng, message) "
                "VALUES (?,?,?,?,?,?)",
                (
                    datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    kind, filename, result.rows_ok, result.rows_ng,
                    " / ".join(result.errors[:3]),
                ),
            )
    return result


def _ensure_referenced_rows(conn, kind: str, rows: list[dict[str, Any]]) -> None:
    """
    来場実績だけ先に取り込んだ場合でも動くように、
    まだ登録されていないコース・顧客を自動で作っておきます。
    """
    if kind in ("visits", "reservations", "budgets"):
        course_ids = {r["course_id"] for r in rows if r.get("course_id")}
        existing = {r[0] for r in conn.execute("SELECT course_id FROM courses")}
        for cid in course_ids - existing:
            conn.execute(
                "INSERT INTO courses(course_id, name) VALUES (?, ?)",
                (cid, f"コース{cid}"),
            )
    if kind in ("visits", "reservations"):
        customer_ids = {r["customer_id"] for r in rows if r.get("customer_id")}
        existing = {r[0] for r in conn.execute("SELECT customer_id FROM customers")}
        for cid in customer_ids - existing:
            conn.execute("INSERT INTO customers(customer_id) VALUES (?)", (cid,))


def import_file(kind: str, content: bytes, filename: str) -> ImportResult:
    """画面からアップロードされたファイルを取り込む入口です。"""
    try:
        rows = read_table(content, filename)
    except Exception as exc:
        res = ImportResult(kind=kind, filename=filename)
        res.errors.append(f"ファイルを開けませんでした: {exc}")
        return res
    return import_rows(kind, rows, filename)
